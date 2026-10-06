from pathlib import Path

import openpyxl
import pytest

from quotes.infrastructure.excel_price_sheet import SheetError, read_price_rows


def workbook(path: Path, data_rows, sheet="PRECIOS GENERAL", header_row=3, extra_empty=0):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = sheet
    ws.cell(row=1, column=3, value="LISTA")
    ws.cell(row=header_row, column=2, value="ITEM")
    ws.cell(row=header_row, column=3, value="NOMBRE")
    ws.cell(row=header_row, column=4, value="PRECIO( SOLES)")
    for offset, (name, price) in enumerate(data_rows, start=1):
        ws.cell(row=header_row + offset, column=3, value=name)
        ws.cell(row=header_row + offset, column=4, value=price)
    last = header_row + len(data_rows) + extra_empty
    for r in range(header_row + len(data_rows) + 1, last + 1):
        ws.cell(row=r, column=3).number_format = "General"  # formatted but empty
    wb.save(path)
    return path


def test_reads_rows_skips_blank_separators_and_trailing_empties(tmp_path):
    path = workbook(
        tmp_path / "w.xlsx",
        [("  TICKETS  CITY TOUR ", 12.34), (None, None), ("GUIA OFICIAL X", 55.5)],
        extra_empty=500,
    )
    result = read_price_rows(path)
    assert [(r.name, str(r.price), r.row) for r in result.rows] == [
        ("TICKETS CITY TOUR", "12.34", 4),
        ("GUIA OFICIAL X", "55.50", 6),
    ]
    assert result.skipped_rows == ()


def test_rows_without_a_usable_price_are_skipped_and_reported(tmp_path):
    path = workbook(
        tmp_path / "w.xlsx", [("A", None), ("B", "n/a"), (None, 5), ("C", -1), ("D", 3)]
    )
    result = read_price_rows(path)
    assert [r.name for r in result.rows] == ["D"]
    assert result.skipped_rows == (4, 5, 6, 7)


def test_header_is_located_even_when_not_on_row_three(tmp_path):
    path = workbook(tmp_path / "w.xlsx", [("A", 1.5)], header_row=5)
    assert [r.row for r in read_price_rows(path).rows] == [6]


def test_custom_sheet_name(tmp_path):
    path = workbook(tmp_path / "w.xlsx", [("A", 1.5)], sheet="OTRA")
    assert len(read_price_rows(path, "OTRA").rows) == 1


def test_missing_file(tmp_path):
    with pytest.raises(SheetError, match="not found"):
        read_price_rows(tmp_path / "nope.xlsx")


def test_not_a_workbook(tmp_path):
    bad = tmp_path / "bad.xlsx"
    bad.write_text("not a workbook")
    with pytest.raises(SheetError, match="Cannot read"):
        read_price_rows(bad)


def test_missing_sheet_lists_available(tmp_path):
    path = workbook(tmp_path / "w.xlsx", [("A", 1)], sheet="OTRA")
    with pytest.raises(SheetError, match="PRECIOS GENERAL.*OTRA"):
        read_price_rows(path)


def test_missing_header(tmp_path):
    path = tmp_path / "w.xlsx"
    wb = openpyxl.Workbook()
    wb.active.title = "PRECIOS GENERAL"
    wb.active["C3"] = "ALGO"
    wb.save(path)
    with pytest.raises(SheetError, match="Header"):
        read_price_rows(path)
