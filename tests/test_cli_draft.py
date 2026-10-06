import io

import openpyxl

from quotes.cli import main
from tests.support.draft_fixture import without_path_lines


def workbook(path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "PRECIOS GENERAL"
    ws["C3"], ws["D3"] = "NOMBRE", "PRECIO( SOLES)"
    for i, (name, price) in enumerate(
        [
            ("TICKETS  INGRESO COLCA", 7310.25),
            ("GUIA OFICIAL  CITY TOUR", 5528.5),
            (None, None),
            ("DESAYUNO COLCA BUFFET", 4391.46),
        ],
        start=4,
    ):
        ws.cell(row=i, column=3, value=name)
        ws.cell(row=i, column=4, value=price)
    wb.save(path)


def run(argv):
    out, err = io.StringIO(), io.StringIO()
    return main(argv, out=out, err=err), out.getvalue(), err.getvalue()


def test_draft_writes_file_and_prints_counts_but_no_prices(tmp_path):
    src, dest = tmp_path / "w.xlsx", tmp_path / "d" / "draft.yaml"
    workbook(src)
    code, out, err = run(["catalog", "draft", str(src), "-o", str(dest)])
    assert code == 0, err
    assert dest.exists()
    assert "pricing items: 2" in out and "local payments: 1" in out
    assert str(dest) in out
    shown = without_path_lines(out + err, src, dest)
    assert "pricing items: 2" in shown  # the filter keeps the summary lines
    for price in ("7310", "5528", "4391"):
        assert price not in shown


def test_draft_refuses_overwrite_unless_forced(tmp_path):
    src, dest = tmp_path / "w.xlsx", tmp_path / "draft.yaml"
    workbook(src)
    assert run(["catalog", "draft", str(src), "-o", str(dest)])[0] == 0
    code, _, err = run(["catalog", "draft", str(src), "-o", str(dest)])
    assert code == 2 and "--force" in err
    assert run(["catalog", "draft", str(src), "-o", str(dest), "--force"])[0] == 0


def test_draft_missing_file_and_sheet_exit_two(tmp_path):
    code, _, err = run(["catalog", "draft", str(tmp_path / "x.xlsx"), "-o", str(tmp_path / "o")])
    assert code == 2 and "not found" in err
    src = tmp_path / "w.xlsx"
    workbook(src)
    code, _, err = run(["catalog", "draft", str(src), "--sheet", "NOPE", "-o", str(tmp_path / "o")])
    assert code == 2 and "NOPE" in err


_BAD_SHEET = (
    b'<?xml version="1.0"?><worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/'
    b'main"><dimension ref="A1:D5"/><sheetData><row r="1"><c r="C1" t="n"><v>1</v></c></row>'
    b"<row r=2><c"
)


def test_draft_with_a_damaged_sheet_exits_two_with_one_line(tmp_path):
    import zipfile

    src = tmp_path / "w.xlsx"
    workbook(src)
    rewritten = tmp_path / "x.tmp"
    with zipfile.ZipFile(src) as source, zipfile.ZipFile(rewritten, "w") as target:
        for item in source.infolist():
            bad = item.filename == "xl/worksheets/sheet1.xml"
            target.writestr(item, _BAD_SHEET if bad else source.read(item))
    rewritten.replace(src)
    code, _, err = run(["catalog", "draft", str(src), "-o", str(tmp_path / "o.yaml")])
    assert code == 2 and err.strip().count("\n") == 0 and "Traceback" not in err
    assert not (tmp_path / "o.yaml").exists()
