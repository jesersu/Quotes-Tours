"""Read the owner's price sheet (read-only, cached values) into plain rows."""

from __future__ import annotations

import zipfile
import zlib
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation
from pathlib import Path
from typing import Any, NamedTuple

import openpyxl
from openpyxl.utils.exceptions import InvalidFileException

from quotes.application.catalog_proposal import RawRow

DEFAULT_SHEET = "PRECIOS GENERAL"
_HEADER_SEARCH_ROWS = 20
_MAX_PRICE = Decimal("99999999.99")  # numeric(10,2)


# In read-only mode openpyxl parses the sheet XML lazily, so a workbook that opens can still
# fail while the header is searched or rows are iterated. These are the families the reader,
# zipfile, zlib and the XML parser raise for damaged input: malformed XML (ParseError is a
# SyntaxError), unreadable or truncated archives and parts, bad numbers or encodings
# (ValueError), and dangling references (LookupError: KeyError/IndexError). Anything else
# (for example a TypeError from our own code) is a bug and is not swallowed.
_READER_FAILURES = (
    OSError,
    EOFError,
    ValueError,
    LookupError,
    SyntaxError,
    zipfile.BadZipFile,
    zlib.error,
    InvalidFileException,
)


class SheetError(Exception):
    """The workbook, sheet or header cannot be used; the message says what to fix."""


class SheetRead(NamedTuple):
    rows: tuple[RawRow, ...]
    skipped_rows: tuple[int, ...]  # rows with a name or a price but not both, or a bad price


def _text(value: Any) -> str:
    return " ".join(str(value).split()) if value is not None else ""


def _price(value: Any) -> Decimal | None:
    if value is None or isinstance(value, bool):
        return None
    try:
        amount = Decimal(str(value).strip())
    except (InvalidOperation, ValueError):  # ValueError: int with thousands of digits
        return None
    # Bounds first: quantizing a huge value (1e30) would raise InvalidOperation.
    if not amount.is_finite() or amount < 0 or amount > _MAX_PRICE:
        return None
    amount = amount.quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)
    return amount if amount <= _MAX_PRICE else None  # rounding can only cross the bound upward


def _find_header(ws: Any) -> int:
    rows = ws.iter_rows(min_row=1, max_row=_HEADER_SEARCH_ROWS, min_col=3, max_col=4)
    for number, (name, price) in enumerate(rows, start=1):
        if _text(name.value).upper() == "NOMBRE" and _text(price.value).upper().startswith(
            "PRECIO"
        ):
            return number
    raise SheetError(
        "Header not found: expected NOMBRE in column C and PRECIO in column D "
        f"within the first {_HEADER_SEARCH_ROWS} rows"
    )


def read_price_rows(path: Path, sheet: str = DEFAULT_SHEET) -> SheetRead:
    path = Path(path)
    if not path.is_file():
        raise SheetError(f"Workbook not found: {path}")
    try:
        workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    except Exception as exc:  # corruption surfaces as several library-specific types
        raise SheetError(f"Cannot read workbook {path}: {type(exc).__name__}") from exc
    try:
        if sheet not in workbook.sheetnames:
            raise SheetError(
                f"Sheet '{sheet}' not found; available sheets: {', '.join(workbook.sheetnames)}"
            )
        ws = workbook[sheet]
        header = _find_header(ws)
        rows: list[RawRow] = []
        skipped: list[int] = []
        for number, (name_cell, price_cell) in enumerate(
            ws.iter_rows(min_row=header + 1, min_col=3, max_col=4, values_only=True),
            start=header + 1,
        ):
            name = _text(name_cell)
            if not name and price_cell is None:
                continue  # blank separator or trailing formatted-but-empty row
            price = _price(price_cell)
            if not name or price is None:
                skipped.append(number)
                continue
            rows.append(RawRow(name, price, number))
        return SheetRead(tuple(rows), tuple(skipped))
    except _READER_FAILURES as exc:
        raise SheetError(f"Cannot read workbook {path}: {type(exc).__name__}") from exc
    finally:
        workbook.close()
